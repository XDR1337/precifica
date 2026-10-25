# -*- coding: utf-8 -*-
"""Importacao da Planilha de Precificacao EG3 (.xlsm/.xlsx) para o banco do aplicativo.

Le os valores digitados nas abas CadInsumos, CadServicos, CadRevenda, CadFopag, CapProd, CF,
FichaTecnica, TblPrecificacao, TblPrecificacaoRevenda, Simulacao, SimulacaoRevenda e Concorrentes.
Os registros sao mesclados pelo nome (nao duplica o que ja existe)."""
from __future__ import annotations

import json

from .core import fnum


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def importar_planilha(db, path: str) -> dict:
    from openpyxl import load_workbook  # import tardio: so carrega quando usado

    wb = load_workbook(path, data_only=True, keep_vba=False)
    names = set(wb.sheetnames)
    rep = {"insumos": 0, "equipe": 0, "servicos": 0, "fichas": 0, "revenda": 0, "custos_fixos": 0,
           "imobilizado": 0, "parametros": 0, "concorrentes": 0}

    def cell(ws, ref):
        v = ws[ref].value
        return None if _is_formula(v) else v

    # ---------------------------------------------------- insumos
    ins_by_name = {r["descricao"].strip().lower(): r for r in db.all("insumos")}
    if "CadInsumos" in names:
        ws = wb["CadInsumos"]
        for r in range(7, ws.max_row + 1):
            desc = _s(ws.cell(r, 4).value)
            custo = ws.cell(r, 8).value
            if not desc or custo in (None, "") or _is_formula(custo):
                continue
            row = dict(descricao=desc, qtde_bruta=fnum(ws.cell(r, 5).value, 1) or 1,
                       qtde_liquida=fnum(ws.cell(r, 6).value, 1) or 1,
                       unidade=_s(ws.cell(r, 7).value).replace("km/h", "Quilômetro") or "Unidade",
                       custo=fnum(custo))
            old = ins_by_name.get(desc.lower())
            if old:
                row["id"] = old["id"]
            else:
                row["categoria"] = "Importado da planilha"
            row["id"] = db.save("insumos", row)
            ins_by_name[desc.lower()] = row
            rep["insumos"] += 1

    # ---------------------------------------------------- equipe + encargos
    emp_by_name = {r["nome"].strip().lower(): r for r in db.all("equipe")}
    if "CadFopag" in names:
        ws = wb["CadFopag"]
        enc = {"enc_ferias": "D9", "enc_13": "D10", "enc_13_ferias": "D11", "enc_inss": "D15", "enc_sat": "D16",
               "enc_sal_educacao": "D17", "enc_sistema_s": "D18", "enc_fgts": "D19", "enc_fgts_rescisao": "D20"}
        vals = {}
        for k, ref in enc.items():
            v = cell(ws, ref)
            if v not in (None, ""):
                vals[k] = round(fnum(v) * 100, 4)
        if vals:
            db.set_many(vals)
            rep["parametros"] += len(vals)
        for r in range(8, 108):
            nome = _s(ws.cell(r, 7).value)
            if not nome:
                continue
            row = dict(nome=nome, salario=fnum(ws.cell(r, 8).value), setor=_s(ws.cell(r, 9).value) or "Produtivo",
                       cargo=_s(ws.cell(r, 10).value), contratacao=_s(ws.cell(r, 11).value) or "CLT",
                       vale_transporte=fnum(ws.cell(r, 14).value), vale_refeicao=fnum(ws.cell(r, 15).value),
                       convenio=fnum(ws.cell(r, 16).value), carga_horaria=fnum(ws.cell(r, 18).value, 176) or 176)
            old = emp_by_name.get(nome.lower())
            if old:
                row["id"] = old["id"]
            row["id"] = db.save("equipe", row)
            emp_by_name[nome.lower()] = row
            rep["equipe"] += 1

    # ---------------------------------------------------- capacidade produtiva
    if "CapProd" in names:
        ws = wb["CapProd"]
        for key, ref, mult in (("horas_dia", "E9", 1), ("dias_mes", "E10", 1), ("produtividade_pct", "E12", 100)):
            v = cell(ws, ref)
            if v not in (None, ""):
                db.set_cfg(key, round(fnum(v) * mult, 4))
                rep["parametros"] += 1

    # ---------------------------------------------------- custos fixos e imobilizado
    if "CF" in names:
        ws = wb["CF"]
        cf_by = {r["descricao"].strip().lower(): r for r in db.all("custos_fixos")}
        for r in range(12, 121):
            desc, val = _s(ws.cell(r, 3).value), ws.cell(r, 5).value
            if not desc or val in (None, "") or _is_formula(val):
                continue
            row = dict(descricao=desc, valor=fnum(val))
            old = cf_by.get(desc.lower())
            if old:
                row["id"] = old["id"]
            db.save("custos_fixos", row)
            rep["custos_fixos"] += 1
        im_by = {r["descricao"].strip().lower(): r for r in db.all("imobilizado")}
        for r in range(10, 27):
            desc, val, vida = _s(ws.cell(r, 7).value), ws.cell(r, 9).value, ws.cell(r, 11).value
            if not desc or val in (None, "") or _is_formula(val):
                continue
            row = dict(descricao=desc, valor=fnum(val), vida_util=fnum(vida, 10) or 10)
            old = im_by.get(desc.lower())
            if old:
                row["id"] = old["id"]
            db.save("imobilizado", row)
            rep["imobilizado"] += 1

    # ---------------------------------------------------- servicos + ficha tecnica + precos + previsao
    srv_by = {r["nome"].strip().lower(): r for r in db.all("servicos")}
    if "CadServicos" in names:
        ws = wb["CadServicos"]
        tp = wb["TblPrecificacao"] if "TblPrecificacao" in names else None
        sim = wb["Simulacao"] if "Simulacao" in names else None
        ft = wb["FichaTecnica"] if "FichaTecnica" in names else None
        prev = {}
        if sim is not None:
            for r in range(8, 108):
                code, q = sim.cell(r, 3).value, sim.cell(r, 6).value
                code = r - 7 if code in (None, "") or _is_formula(code) else int(fnum(code))
                if q not in (None, "") and not _is_formula(q):
                    prev[code] = fnum(q)
        for n in range(1, 101):
            nome = _s(ws.cell(6 + n, 4).value)
            if not nome:
                continue
            row = dict(nome=nome, previsao=prev.get(n, 0))
            if tp is not None:
                r = 6 + n
                row.update(preco=fnum(tp.cell(r, 10).value), imposto_pct=fnum(tp.cell(r, 14).value) * 100,
                           cartao_pct=fnum(tp.cell(r, 16).value) * 100, app_pct=fnum(tp.cell(r, 18).value) * 100,
                           comissao_pct=fnum(tp.cell(r, 20).value) * 100)
            old = srv_by.get(nome.lower())
            if old:
                row["id"] = old["id"]
            sid = db.save("servicos", row)
            srv_by[nome.lower()] = dict(row, id=sid)
            rep["servicos"] += 1
            if ft is None:
                continue
            idx = n - 1
            top = 6 + 35 * (idx // 2)
            c_name, c_qty = (3, 6) if idx % 2 == 0 else (9, 12)
            mod, ins = [], []
            for r in range(top + 2, top + 10):
                nm, h = _s(ft.cell(r, c_name).value), ft.cell(r, c_qty).value
                if nm and h not in (None, "") and nm.lower() in emp_by_name:
                    mod.append((emp_by_name[nm.lower()]["id"], fnum(h)))
            for r in range(top + 11, top + 30):
                nm, q = _s(ft.cell(r, c_name).value), ft.cell(r, c_qty).value
                if nm and q not in (None, "") and nm.lower() in ins_by_name:
                    ins.append((ins_by_name[nm.lower()]["id"], fnum(q)))
            if mod or ins:
                db.execute("DELETE FROM ficha_mod WHERE servico_id = ?", (sid,))
                db.execute("DELETE FROM ficha_insumo WHERE servico_id = ?", (sid,))
                for eid, h in mod:
                    db.save("ficha_mod", dict(servico_id=sid, equipe_id=eid, horas=h))
                for iid, q in ins:
                    db.save("ficha_insumo", dict(servico_id=sid, insumo_id=iid, quantidade=q))
                rep["fichas"] += 1

    # ---------------------------------------------------- revenda
    if "CadRevenda" in names:
        ws = wb["CadRevenda"]
        tp = wb["TblPrecificacaoRevenda"] if "TblPrecificacaoRevenda" in names else None
        sim = wb["SimulacaoRevenda"] if "SimulacaoRevenda" in names else None
        prev = {}
        if sim is not None:
            for r in range(8, 1008):
                code, q = sim.cell(r, 3).value, sim.cell(r, 6).value
                code = r - 7 if code in (None, "") or _is_formula(code) else int(fnum(code))
                if q not in (None, "") and not _is_formula(q):
                    prev[code] = fnum(q)
        rev_by = {r["nome"].strip().lower(): r for r in db.all("revenda")}
        for r in range(7, ws.max_row + 1):
            nome = _s(ws.cell(r, 4).value)
            if not nome:
                continue
            code = int(fnum(ws.cell(r, 3).value, r - 6)) or (r - 6)
            row = dict(nome=nome, unidade=_s(ws.cell(r, 5).value) or "Unidade", qtde=fnum(ws.cell(r, 6).value, 1) or 1,
                       custo=fnum(ws.cell(r, 7).value), previsao=prev.get(code, 0))
            if tp is not None:
                tr = 6 + code
                row.update(preco=fnum(tp.cell(tr, 10).value), imposto_pct=fnum(tp.cell(tr, 12).value) * 100,
                           cartao_pct=fnum(tp.cell(tr, 14).value) * 100, app_pct=fnum(tp.cell(tr, 16).value) * 100,
                           comissao_pct=fnum(tp.cell(tr, 18).value) * 100)
            old = rev_by.get(nome.lower())
            if old:
                row["id"] = old["id"]
            db.save("revenda", row)
            rep["revenda"] += 1

    # ---------------------------------------------------- lucro desejado
    if "Analise" in names:
        v = cell(wb["Analise"], "R9")
        if v not in (None, ""):
            db.set_cfg("lucro_desejado_pct", round(fnum(v) * 100, 4))

    # ---------------------------------------------------- concorrentes
    if "Concorrentes" in names:
        ws = wb["Concorrentes"]
        itens = [_s(ws.cell(r, 3).value).split(". ", 1)[-1] for r in range(16, 22)]
        empresas = [_s(ws.cell(15, c).value).title() or f"Empresa {c}" for c in range(4, 8)]
        notas, precos, alvo = {}, {}, []
        for i, r in enumerate(range(16, 22)):
            for j, c in enumerate(range(4, 8)):
                v = ws.cell(r, c).value
                if v not in (None, "") and not _is_formula(v):
                    notas[f"{i}|{j}"] = fnum(v)
        for j, c in enumerate(range(4, 8)):
            v = ws.cell(23, c).value
            if v not in (None, "") and not _is_formula(v):
                precos[str(j)] = fnum(v)
        for r in range(7, 13):
            for c in range(3, 8):
                v = _s(ws.cell(r, c).value)
                if v and not v.startswith("="):
                    alvo.append(v)
        if notas or precos or alvo:
            dados = json.loads(db.cfg("concorrentes_json") or "{}") if db.cfg("concorrentes_json") else {}
            dados.update(empresas=empresas, itens=itens, notas=notas, precos=precos)
            if alvo:
                dados["cliente_alvo"] = "\n".join(alvo)
            db.set_cfg("concorrentes_json", json.dumps(dados, ensure_ascii=False))
            rep["concorrentes"] = len(notas) + len(precos)
    return rep
