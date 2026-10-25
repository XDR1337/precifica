"""Monta uma planilha EG3 sintetica (celulas nas mesmas posicoes que o importador espera) e importa com o importador Python.
Saidas em test/out/: planilha_teste.xlsx, import_expected.json (tabelas depois da importacao)."""
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(ROOT))

from precifica.db import Database, SCHEMA      # noqa: E402
from precifica.importer import importar_planilha  # noqa: E402

OUT = HERE / "out"
OUT.mkdir(exist_ok=True)
Database.migrate_old_csv = lambda self: 0

wb = Workbook()
wb.remove(wb.active)


def sheet(name):
    return wb.create_sheet(name)


ws = sheet("CadInsumos")
rows = [("Cabo flexível 2,5 mm", 100, 95, "Metro", 250), ("Disjuntor 20A (cx)", 10, 10, "Unidade", 120), ("Deslocamento", 1, 1, "km/h", 1.9), ("Fita isolante", 20, 18, "Rolo", 60)]
for i, (d, b, l, u, c) in enumerate(rows):
    r = 7 + i
    ws.cell(r, 4, d); ws.cell(r, 5, b); ws.cell(r, 6, l); ws.cell(r, 7, u); ws.cell(r, 8, c)
ws.cell(12, 4, "Sem custo"); ws.cell(12, 8, None)

ws = sheet("CadFopag")
for ref, v in dict(D9=0.1111, D10=0.0833, D11=0.0233, D15=0, D16=0.01, D17=0, D18=0, D19=0.08, D20=0.04).items():
    ws[ref] = v
emps = [("Marcos", 3500, "Produtivo", "Eletricista", "CLT", 200, 400, 150, 176), ("Paula", 5000, "Administrativo", "Gerente", "Pró-Labore", 0, 0, 0, 160), ("Rui", 2000, "Produtivo", "Aux", "PJ", 0, 0, 0, 0)]
for i, (n, s, st, c, ct, vt, vr, cv, h) in enumerate(emps):
    r = 8 + i
    ws.cell(r, 7, n); ws.cell(r, 8, s); ws.cell(r, 9, st); ws.cell(r, 10, c); ws.cell(r, 11, ct); ws.cell(r, 14, vt); ws.cell(r, 15, vr); ws.cell(r, 16, cv); ws.cell(r, 18, h)

ws = sheet("CapProd")
ws["E9"], ws["E10"], ws["E12"] = 8, 21, 0.7

ws = sheet("CF")
for i, (d, v) in enumerate([("Aluguel", 1800), ("Contador", 450.5), ("Internet", 120)]):
    ws.cell(12 + i, 3, d); ws.cell(12 + i, 5, v)
for i, (d, v, vida) in enumerate([("Máquinas e Equipamentos", 8000, 10), ("Veículos", 45000, 5), ("Edifícios", 0, 25)]):
    ws.cell(10 + i, 7, d); ws.cell(10 + i, 9, v); ws.cell(10 + i, 11, vida)

serv = ["Instalação de tomada", "Troca de disjuntor", "Quadro elétrico", "Aterramento", "Serviço sem ficha"]
ws = sheet("CadServicos")
for n, nome in enumerate(serv, 1):
    ws.cell(6 + n, 4, nome)
tp = sheet("TblPrecificacao")
for n, (p, imp, car, app, com) in enumerate([(100, .06, .03, 0, 0), (80, .06, .03, .02, .01), (900, .06, 0, 0, .05), (300, .06, .025, 0, 0), (50, .06, 0, 0, 0)], 1):
    r = 6 + n
    tp.cell(r, 10, p); tp.cell(r, 14, imp); tp.cell(r, 16, car); tp.cell(r, 18, app); tp.cell(r, 20, com)
sim = sheet("Simulacao")
for n, q in enumerate([30, 12, 4, 0, 6], 1):
    sim.cell(7 + n, 3, n); sim.cell(7 + n, 6, q)
ft = sheet("FichaTecnica")
fichas = {1: ([("Marcos", 1.0)], [("Cabo flexível 2,5 mm", 5), ("Deslocamento", 12)]), 2: ([("Marcos", 0.5), ("Rui", 0.5)], [("Disjuntor 20A (cx)", 1)]),
          3: ([("Marcos", 4)], [("Fita isolante", 2), ("Deslocamento", 20)]), 4: ([("Rui", 2)], [])}
for n, (mod, ins) in fichas.items():
    idx = n - 1
    top = 6 + 35 * (idx // 2)
    cn, cq = (3, 6) if idx % 2 == 0 else (9, 12)
    for k, (nm, h) in enumerate(mod):
        ft.cell(top + 2 + k, cn, nm); ft.cell(top + 2 + k, cq, h)
    for k, (nm, q) in enumerate(ins):
        ft.cell(top + 11 + k, cn, nm); ft.cell(top + 11 + k, cq, q)

ws = sheet("CadRevenda")
rev = [(1, "Tomada 10A", "Unidade", 10, 80), (2, "Lâmpada LED", "Unidade", 20, 150), (3, "Fita LED 5m", "Rolo", 1, 35)]
for i, (c, n, u, q, v) in enumerate(rev):
    r = 7 + i
    ws.cell(r, 3, c); ws.cell(r, 4, n); ws.cell(r, 5, u); ws.cell(r, 6, q); ws.cell(r, 7, v)
tr = sheet("TblPrecificacaoRevenda")
for c, (p, i, ca, a, co) in enumerate([(15, .06, .03, 0, 0), (12, .06, 0, 0, .02), (70, .06, .025, .01, 0)], 1):
    r = 6 + c
    tr.cell(r, 10, p); tr.cell(r, 12, i); tr.cell(r, 14, ca); tr.cell(r, 16, a); tr.cell(r, 18, co)
sr = sheet("SimulacaoRevenda")
for c, q in enumerate([50, 0, 7], 1):
    sr.cell(7 + c, 3, c); sr.cell(7 + c, 6, q)

sheet("Analise")["R9"] = 0.22
ws = sheet("Concorrentes")
for i, t in enumerate(["1. Atendimento", "2. Qualidade", "3. Pós-venda", "4. Localização", "5. Marca", "6. Pagamento"]):
    ws.cell(16 + i, 3, t)
for j, e in enumerate(["EMPRESA ALFA", "beta ltda", "Gama", "NOSSA EMPRESA"]):
    ws.cell(15, 4 + j, e)
for i in range(6):
    for j in range(4):
        if (i + j) % 3:
            ws.cell(16 + i, 4 + j, (i + j) % 3 + 1)
for j, p in enumerate([200, 180, 0, 190]):
    if p:
        ws.cell(23, 4 + j, p)
ws.cell(7, 3, "Casas de médio padrão"); ws.cell(8, 4, "Renda 5-10 SM")

xlsx = OUT / "planilha_teste.xlsx"
wb.save(xlsx)

db = Database(Path(tempfile.mkdtemp()) / "t.db")
rep = importar_planilha(db, str(xlsx))
state = {t: db.all(t) for t in SCHEMA if t != "config"}
state["config"] = db.config()
(OUT / "import_expected.json").write_text(json.dumps(dict(rep=rep, state=state), ensure_ascii=False), encoding="utf-8")
print("import python:", rep)
